"""Flats and members import from Excel (docs/02 §4).

Never import silently: upload → preview with a status per row → confirm. Confirm re-checks
every row against the database (someone may have added a flat since the preview) and imports
OK and warning rows in one transaction. Error rows are skipped and listed.
"""

import io
import zipfile
from dataclasses import asdict
from typing import Any
from uuid import UUID, uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import get_settings
from app.errors import AppError, not_found
from app.models import Attachment, Flat, ImportJob, Member, User, Wing
from app.services import audit, ledger
from app.services.import_rows import (
    COLUMNS,
    Issue,
    RowResult,
    check_rows,
    normalise_header,
)
from app.services.societies import DEFAULT_WING_NAME
from app.storage import store
from app.timeutil import today_ist

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEET_NAME = "Flats"
MAX_ROWS = 5000

INSTRUCTIONS = [
    "How to fill the Flats sheet",
    "One row per flat. Keep the column headings as they are.",
    "Wing: the building or wing name, e.g. A. Leave empty only if the society has no wings.",
    "Wing has lift: Y or N. Lift charges are shared only by flats in wings with a lift.",
    "Flat number: exactly as the society writes it, e.g. B-207 or 301.",
    "Carpet area: as in the registered agreement, in square feet.",
    "Usage: Residential or Commercial.",
    "Certified construction cost: the architect-certified cost of the flat, in rupees.",
    "Water taps: number of taps or inlets as per the sanctioned plan.",
    "Occupancy: Owner occupied, Family occupied, Let out or Vacant.",
    "Phones: 10 digits. Email is optional.",
    "Opening balance: amount the flat owes today, in rupees. Use a minus sign for an advance.",
    "Nothing is saved until you check the preview and confirm.",
]


def build_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = SHEET_NAME
    sheet.append([header for _, header in COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for index, (_, header) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[cell_letter(index)].width = max(14, len(header) + 2)
    sheet.freeze_panes = "A2"

    choices = {
        "wing_has_lift": '"Y,N"',
        "usage": '"Residential,Commercial"',
        "occupancy": '"Owner occupied,Family occupied,Let out,Vacant"',
    }
    for index, (key, _) in enumerate(COLUMNS, start=1):
        if key in choices:
            rule = DataValidation(type="list", formula1=choices[key], allow_blank=True)
            sheet.add_data_validation(rule)
            rule.add(f"{cell_letter(index)}2:{cell_letter(index)}{MAX_ROWS + 1}")

    help_sheet = workbook.create_sheet("How to fill")
    for line in INSTRUCTIONS:
        help_sheet.append([line])
    help_sheet["A1"].font = Font(bold=True)
    help_sheet.column_dimensions["A"].width = 100

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def cell_letter(index: int) -> str:
    letters = ""
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def read_workbook(data: bytes) -> list[tuple[int, dict[str, object]]]:
    """Rows keyed by column key. Raises AppError for files that are not our template."""
    if not data.startswith(b"PK\x03\x04"):
        raise AppError(422, "NOT_XLSX", "errors.notXlsx")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise AppError(422, "NOT_XLSX", "errors.notXlsx") from exc
    sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.worksheets[0]

    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None) or ()
    by_header = {normalise_header(h): key for key, h in COLUMNS}
    positions = {
        by_header[normalise_header(text)]: i
        for i, text in enumerate(header)
        if normalise_header(text) in by_header
    }
    if "flat_number" not in positions:
        raise AppError(422, "TEMPLATE_COLUMNS", "errors.templateColumns")

    result: list[tuple[int, dict[str, object]]] = []
    for row_number, values in enumerate(rows, start=2):
        raw = {key: values[i] if i < len(values) else None for key, i in positions.items()}
        if all(v is None or str(v).strip() == "" for v in raw.values()):
            continue
        result.append((row_number, raw))
        if len(result) > MAX_ROWS:
            raise AppError(422, "TOO_MANY_ROWS", "errors.tooManyRows")
    workbook.close()
    return result


async def _existing(
    session: AsyncSession, society_id: UUID
) -> tuple[set[tuple[str, str]], dict[str, bool], str]:
    """Saved flats as (wing, number) lower-cased, lift by wing, and the wing that rows
    without a wing name go into (the only wing, else the default one)."""
    rows = await session.execute(
        select(Wing.name, Flat.flat_number)
        .join(Flat, Flat.wing_id == Wing.id)
        .where(Flat.society_id == society_id)
    )
    flats = {(wing.lower(), number.lower()) for wing, number in rows}
    wings = list(await session.scalars(select(Wing).where(Wing.society_id == society_id)))
    blank = wings[0].name if len(wings) == 1 else DEFAULT_WING_NAME
    return flats, {w.name.lower(): w.has_lift for w in wings}, blank.lower()


def _row_json(result: RowResult) -> dict[str, Any]:
    return {
        "row": result.row,
        "status": result.status,
        "issues": [asdict(i) for i in result.issues],
        "data": result.data,
    }


def _summary(results: list[RowResult]) -> dict[str, int]:
    return {s: sum(1 for r in results if r.status == s) for s in ("ok", "warning", "error")}


async def preview(
    session: AsyncSession, society_id: UUID, actor: User, filename: str | None, data: bytes
) -> ImportJob:
    if len(data) > get_settings().max_upload_bytes:
        raise AppError(413, "FILE_TOO_LARGE", "errors.fileTooLarge")
    raw_rows = read_workbook(data)
    stored = await store(society_id, "imports", "xlsx", data, XLSX_MIME)
    # Attachments are insert-only, so link the file to the job id before either is saved.
    job_id = uuid4()
    attachment = Attachment(
        society_id=society_id,
        owner_type="import_job",
        owner_id=job_id,
        storage_path=stored.path,
        mime=XLSX_MIME,
        size_bytes=stored.size_bytes,
        sha256=stored.sha256,
        original_filename=(filename or "")[:200] or None,
        uploaded_by=actor.id,
    )
    session.add(attachment)
    await session.flush()

    flats, lifts, blank_wing = await _existing(session, society_id)
    results = check_rows(raw_rows, flats, lifts, blank_wing)
    summary = _summary(results)
    job = ImportJob(
        id=job_id,
        society_id=society_id,
        kind="flats_members",
        attachment_id=attachment.id,
        status="previewed",
        row_count=len(results),
        ok_count=summary["ok"] + summary["warning"],
        error_count=summary["error"],
        preview={"rows": [_row_json(r) for r in results], "summary": summary},
        created_by=actor.id,
    )
    session.add(job)
    await session.flush()
    await audit.record(
        session,
        society_id=society_id,
        actor_user_id=actor.id,
        action="import.preview",
        entity_type="import_job",
        entity_id=job.id,
        after={"rows": len(results), **summary, "file_sha256": stored.sha256},
    )
    await session.refresh(job)
    return job


async def get_job(session: AsyncSession, society_id: UUID, job_id: UUID) -> ImportJob:
    job = await session.get(ImportJob, job_id)
    if job is None or job.society_id != society_id:
        raise not_found()
    return job


async def _wing_for(
    session: AsyncSession,
    society_id: UUID,
    actor: User,
    name: str | None,
    has_lift: bool | None,
    cache: dict[str, Wing],
) -> Wing:
    wings = cache or {w.name.lower(): w for w in await _all_wings(session, society_id)}
    cache.update(wings)
    if name is None:
        # Whole file without wing names: use the society's only wing, else the default one.
        if len(cache) == 1:
            return next(iter(cache.values()))
        name = DEFAULT_WING_NAME
    wing = cache.get(name.lower())
    if wing is None:
        wing = Wing(society_id=society_id, name=name, has_lift=bool(has_lift), created_by=actor.id)
        session.add(wing)
        await session.flush()
        cache[name.lower()] = wing
        await audit.record(
            session,
            society_id=society_id,
            actor_user_id=actor.id,
            action="wing.create",
            entity_type="wing",
            entity_id=wing.id,
            after={"name": name, "has_lift": wing.has_lift, "source": "import"},
        )
    return wing


async def _all_wings(session: AsyncSession, society_id: UUID) -> list[Wing]:
    return list(await session.scalars(select(Wing).where(Wing.society_id == society_id)))


async def confirm(session: AsyncSession, society_id: UUID, actor: User, job_id: UUID) -> ImportJob:
    job = await get_job(session, society_id, job_id)
    if job.status != "previewed":
        raise AppError(409, "IMPORT_DONE", "errors.importDone")

    # Re-check against the database as it is now; the stored data is already normalised,
    # so only the "already exists" check can change.
    flats, _, blank_wing = await _existing(session, society_id)
    rows = job.preview["rows"]
    entry_date = today_ist()
    wings: dict[str, Wing] = {}
    imported = 0
    for row in rows:
        data = row["data"]
        if row["status"] != "error":
            key = ((data["wing"] or blank_wing).lower(), data["flat_number"].lower())
            if key in flats:
                row["status"] = "error"
                row["issues"].insert(0, asdict(Issue("alreadyExists")))
        if row["status"] == "error":
            continue

        wing = await _wing_for(
            session, society_id, actor, data["wing"], data["wing_has_lift"], wings
        )
        flat = Flat(
            society_id=society_id,
            wing_id=wing.id,
            flat_number=data["flat_number"],
            carpet_area_sqft=data["carpet_area_sqft"],
            flat_type=data["flat_type"],
            usage=data["usage"],
            certified_construction_cost_paise=data["construction_cost_paise"],
            water_points=data["water_points"],
            municipal_property_account=data["municipal_account"],
            occupancy=data["occupancy"],
            created_by=actor.id,
        )
        session.add(flat)
        await session.flush()
        if data["owner_name"]:
            session.add(
                Member(
                    society_id=society_id,
                    flat_id=flat.id,
                    full_name=data["owner_name"],
                    phone_e164=data["owner_phone"],
                    email=data["owner_email"],
                    member_type="owner",
                    is_primary=True,
                    created_by=actor.id,
                )
            )
        if data["tenant_name"]:
            session.add(
                Member(
                    society_id=society_id,
                    flat_id=flat.id,
                    full_name=data["tenant_name"],
                    phone_e164=data["tenant_phone"],
                    member_type="tenant",
                    created_by=actor.id,
                )
            )
        if data["opening_balance_paise"]:
            await ledger.add_opening_balance(
                session,
                society_id=society_id,
                flat_id=flat.id,
                amount_paise=data["opening_balance_paise"],
                entry_date=entry_date,
                source_type="import",
                source_id=job.id,
                created_by=actor.id,
            )
        await session.flush()
        await audit.record(
            session,
            society_id=society_id,
            actor_user_id=actor.id,
            action="flat.import",
            entity_type="flat",
            entity_id=flat.id,
            after={
                "import_job_id": str(job.id),
                "row": row["row"],
                "flat_number": data["flat_number"],
                "wing": wing.name,
                "owner_name": data["owner_name"],
                "tenant_name": data["tenant_name"],
                "opening_balance_paise": data["opening_balance_paise"],
            },
        )
        flats.add((wing.name.lower(), data["flat_number"].lower()))
        imported += 1

    skipped = len(rows) - imported
    job.preview = {**job.preview, "rows": rows, "imported": imported, "skipped": skipped}
    job.status = "imported"
    job.ok_count = imported
    job.error_count = skipped
    await session.flush()
    await audit.record(
        session,
        society_id=society_id,
        actor_user_id=actor.id,
        action="import.confirm",
        entity_type="import_job",
        entity_id=job.id,
        after={"imported": imported, "skipped": skipped},
    )
    await session.refresh(job)
    return job
