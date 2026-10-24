"""Excel templates and reading them back, shared by every import (docs/02 §4, §5)."""

import io
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from app.errors import AppError
from app.services.import_rows import normalise_header

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_ROWS = 5000

Columns = list[tuple[str, str]]


def cell_letter(index: int) -> str:
    letters = ""
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def build_template(
    sheet_name: str, columns: Columns, choices: dict[str, str], instructions: list[str]
) -> bytes:
    """choices: column key → comma-separated allowed values for a dropdown."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = sheet_name
    sheet.append([header for _, header in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for index, (key, header) in enumerate(columns, start=1):
        letter = cell_letter(index)
        sheet.column_dimensions[letter].width = max(14, len(header) + 2)
        if key in choices:
            rule = DataValidation(type="list", formula1=f'"{choices[key]}"', allow_blank=True)
            sheet.add_data_validation(rule)
            rule.add(f"{letter}2:{letter}{MAX_ROWS + 1}")
    sheet.freeze_panes = "A2"

    help_sheet = workbook.create_sheet("How to fill")
    for line in instructions:
        help_sheet.append([line])
    help_sheet["A1"].font = Font(bold=True)
    help_sheet.column_dimensions["A"].width = 100

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def read_workbook(
    data: bytes, sheet_name: str, columns: Columns, required_key: str
) -> list[tuple[int, dict[str, object]]]:
    """Rows keyed by column key, with their Excel row numbers. Raises AppError for files that
    are not our template."""
    if not data.startswith(b"PK\x03\x04"):
        raise AppError(422, "NOT_XLSX", "errors.notXlsx")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise AppError(422, "NOT_XLSX", "errors.notXlsx") from exc
    sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.worksheets[0]

    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None) or ()
    by_header = {normalise_header(h): key for key, h in columns}
    positions = {
        by_header[normalise_header(text)]: i
        for i, text in enumerate(header)
        if normalise_header(text) in by_header
    }
    if required_key not in positions:
        raise AppError(422, "TEMPLATE_COLUMNS", "errors.templateColumns")

    result: list[tuple[int, dict[str, object]]] = []
    for row_number, values in enumerate(rows, start=2):
        raw = {key: values[i] if i < len(values) else None for key, i in positions.items()}
        if all(v is None or str(v).strip() == "" for v in raw.values()):
            continue
        result.append((row_number, raw))
        if len(result) > MAX_ROWS:
            raise AppError(422, "TOO_MANY_ROWS", "errors.tooManyRows")
    workbook.close()
    return result
