"""All of a society's data as Excel and PDF (docs/07 M11, docs/02 §2 "Export data").

Every table with a society_id is exported in full, so nothing is left out when a new table is
added. Rows are read through the normal request session, so row-level security keeps the
export to the current society even if a query here were wrong. Secrets (invite token hashes)
are never exported. Money columns (``*_paise``) are shown as exact rupees (Decimal, never
float) next to the column name with "(₹)".
"""

import io
import json
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from typing import Any
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import Table, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Base, User
from app.services import audit
from app.timeutil import IST

# Never leave the server, even for the society's own admin.
SECRET_COLUMNS = {"invite_token_hash"}
# Excel limits a sheet name to 31 characters.
SHEET_NAME_LIMIT = 31


@dataclass
class Sheet:
    name: str
    headers: list[str]
    rows: list[list[Any]]


def society_tables() -> list[Table]:
    return sorted(
        (t for t in Base.metadata.tables.values() if "society_id" in t.columns),
        key=lambda t: t.name,
    )


def header(column: str) -> str:
    return f"{column.removesuffix('_paise')} (₹)" if column.endswith("_paise") else column


def cell(column: str, value: object) -> object:
    """Excel-safe value: exact rupees for paise, India time for timestamps, text for ids."""
    if value is None:
        return None
    if column.endswith("_paise") and isinstance(value, int):
        return Decimal(value) / 100
    if isinstance(value, datetime):
        return value.astimezone(IST).replace(tzinfo=None) if value.tzinfo else value
    if isinstance(value, UUID):
        return str(value)
    if isinstance(value, dict | list):
        return json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, bool | int | Decimal | date | str):
        return value
    return str(value)


async def collect(session: AsyncSession, society_id: UUID) -> list[Sheet]:
    sheets = []
    for table in society_tables():
        columns = [c for c in table.columns if c.name not in SECRET_COLUMNS]
        order = [table.c.created_at] if "created_at" in table.c else []
        result = await session.execute(
            select(*columns).where(table.c.society_id == society_id).order_by(*order)
        )
        names = [c.name for c in columns]
        sheets.append(
            Sheet(
                name=table.name[:SHEET_NAME_LIMIT],
                headers=[header(n) for n in names],
                rows=[[cell(n, v) for n, v in zip(names, row, strict=True)] for row in result],
            )
        )
    return sheets


def to_xlsx(sheets: list[Sheet]) -> bytes:
    workbook = Workbook()
    default = workbook.active
    if default is not None:
        workbook.remove(default)
    for s in sheets:
        sheet = workbook.create_sheet(s.name)
        sheet.append(s.headers)
        for c in sheet[1]:
            c.font = Font(bold=True)
        for row in s.rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def as_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return f"{value:.2f}"
    if isinstance(value, datetime):
        return value.strftime("%d %b %Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d %b %Y")
    return str(value)


async def record_export(
    session: AsyncSession, society_id: UUID, actor: User, file_format: str, sheets: list[Sheet]
) -> None:
    await audit.record(
        session,
        society_id=society_id,
        actor_user_id=actor.id,
        action="society.export",
        entity_type="society",
        entity_id=society_id,
        after={"format": file_format, "rows": {s.name: len(s.rows) for s in sheets}},
    )
