"""Compare two SocietyIQ Excel exports sheet by sheet (monthly restore test, docs/09 §6c).

Run from backend/: uv run python scripts/compare_exports.py live.xlsx restored.xlsx
Exit code 0 when every sheet has the same rows; 1 otherwise, with the differences listed.
Only row counts and row contents are compared; nothing is printed from the rows themselves,
because they hold members' personal data.
"""

import hashlib
import sys

from openpyxl import load_workbook


def fingerprint(path: str) -> dict[str, tuple[int, str]]:
    book = load_workbook(path, read_only=True)
    out = {}
    for sheet in book.worksheets:
        digest = hashlib.sha256()
        count = 0
        for row in sheet.iter_rows(min_row=2, values_only=True):
            digest.update(repr(row).encode())
            count += 1
        out[sheet.title] = (count, digest.hexdigest())
    return out


def main(live: str, restored: str) -> int:
    a, b = fingerprint(live), fingerprint(restored)
    problems = []
    for name in sorted(set(a) | set(b)):
        if name not in b:
            problems.append(f"{name}: missing in the restored export")
        elif name not in a:
            problems.append(f"{name}: only in the restored export")
        elif a[name][0] != b[name][0]:
            problems.append(f"{name}: {a[name][0]} rows live, {b[name][0]} restored")
        elif a[name][1] != b[name][1]:
            problems.append(f"{name}: same row count, different contents")
    for line in problems:
        print(line)
    total = sum(count for count, _ in a.values())
    print(f"{len(a)} sheets, {total} rows live; {'OK' if not problems else 'DIFFERENT'}")
    return 1 if problems else 0


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: compare_exports.py live.xlsx restored.xlsx")
    raise SystemExit(main(sys.argv[1], sys.argv[2]))
