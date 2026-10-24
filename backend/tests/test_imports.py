"""M3 acceptance (docs/07): importing the template with 3 deliberately wrong rows shows
exactly those 3 errors; confirming imports the rest; ledger balances equal the opening
balances."""

import io

import asyncpg
from openpyxl import Workbook, load_workbook

from app.services.import_rows import COLUMNS
from tests.conftest import APP_URL, needs_db

pytestmark = needs_db

HEADERS = [h for _, h in COLUMNS]


def _sheet(rows: list[dict[str, object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Flats"
    ws.append(HEADERS)
    keys = [k for k, _ in COLUMNS]
    for row in rows:
        ws.append([row.get(k) for k in keys])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _flat(wing: str, number: str, balance: object = None, **extra: object) -> dict[str, object]:
    return {
        "wing": wing,
        "wing_has_lift": "N",
        "flat_number": number,
        "carpet_area_sqft": 650,
        "usage": "Residential",
        "construction_cost": 960000,
        "occupancy": "Owner occupied",
        "owner_name": f"Member {wing}-{number}",
        "opening_balance": balance,
        **extra,
    }


GOOD = [
    _flat("A", "101", 0),
    _flat("A", "204", 4200),
    _flat("A", "305", "12,600", occupancy="Let out", tenant_name="Tenant A-305"),
    _flat("B", "207", 29400, owner_phone="98000 00000"),
    _flat("B", "102", -500),
]
WRONG = [
    _flat("C", "302", 16800, owner_phone="98000000"),  # 8 digits
    _flat("C", None, 100),  # no flat number
    _flat("C", "404", "12.345"),  # fraction of a paisa
]


def _upload(client, society_id, headers, data: bytes):  # noqa: ANN001, ANN202
    return client.post(
        f"/api/v1/societies/{society_id}/imports",
        files={"file": ("flats.xlsx", data, "application/octet-stream")},
        headers=headers,
    )


async def _balances(society_id) -> dict[str, int]:  # noqa: ANN001
    conn = await asyncpg.connect(APP_URL)
    try:
        async with conn.transaction():
            await conn.execute(
                "SELECT set_config('app.current_society_id', $1, true)", str(society_id)
            )
            rows = await conn.fetch(
                """SELECT w.name || '-' || f.flat_number AS flat,
                          coalesce(sum(l.amount_paise), 0) AS balance,
                          count(l.id) AS entries
                   FROM flats f JOIN wings w ON w.id = f.wing_id
                   LEFT JOIN ledger_entries l ON l.flat_id = f.id
                   GROUP BY 1"""
            )
    finally:
        await conn.close()
    return {r["flat"]: (int(r["balance"]), int(r["entries"])) for r in rows}


async def test_template_downloads_with_headers(client, two_societies) -> None:  # noqa: ANN001
    s = two_societies
    r = await client.get(f"/api/v1/societies/{s.a_id}/imports/template", headers=s.a_admin.headers)
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content))
    assert [c.value for c in wb["Flats"][1]] == HEADERS


async def test_three_wrong_rows_show_exactly_three_errors_then_import_the_rest(
    client,  # noqa: ANN001
    two_societies,  # noqa: ANN001
) -> None:
    s = two_societies
    h = s.a_admin.headers
    r = await _upload(client, s.a_id, h, _sheet(GOOD + WRONG))
    assert r.status_code == 201, r.text
    job = r.json()
    errors = [row for row in job["rows"] if row["status"] == "error"]
    assert [row["row"] for row in errors] == [7, 8, 9]
    assert errors[0]["issues"][0] == {
        "key": "phoneDigits",
        "params": {"field": "owner_phone", "digits": 8},
    }
    assert errors[1]["issues"][0]["key"] == "required"
    assert errors[2]["issues"][0]["key"] == "tooManyDecimals"
    assert job["ok_count"] == 5
    assert job["error_count"] == 3

    # Nothing is saved before confirming.
    assert await _balances(s.a_id) == {}

    r = await client.post(f"/api/v1/societies/{s.a_id}/imports/{job['id']}/confirm", headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "imported"
    assert r.json()["ok_count"] == 5

    # Ledger balances equal the opening balances; no entry for a zero balance.
    assert await _balances(s.a_id) == {
        "A-101": (0, 0),
        "A-204": (420000, 1),
        "A-305": (1260000, 1),
        "B-207": (2940000, 1),
        "B-102": (-50000, 1),
    }

    flats = (await client.get(f"/api/v1/societies/{s.a_id}/flats", headers=h)).json()
    by_number = {f"{f['wing_name']}-{f['flat_number']}": f for f in flats}
    assert by_number["A-305"]["occupancy"] == "let_out"
    assert by_number["B-207"]["balance_paise"] == 2940000
    detail = (
        await client.get(f"/api/v1/societies/{s.a_id}/flats/{by_number['A-305']['id']}", headers=h)
    ).json()
    assert sorted(m["member_type"] for m in detail["members"]) == ["owner", "tenant"]

    # Confirming twice is refused; uploading the same file again flags every row.
    again = await client.post(f"/api/v1/societies/{s.a_id}/imports/{job['id']}/confirm", headers=h)
    assert again.status_code == 409
    second = (await _upload(client, s.a_id, h, _sheet(GOOD))).json()
    assert {row["issues"][0]["key"] for row in second["rows"]} == {"alreadyExists"}


async def test_flat_added_after_preview_is_skipped_on_confirm(client, two_societies) -> None:  # noqa: ANN001
    s = two_societies
    h = s.a_admin.headers
    job = (
        await _upload(client, s.a_id, h, _sheet([_flat("Main", "1"), _flat("Main", "2")]))
    ).json()
    wing = (await client.get(f"/api/v1/societies/{s.a_id}/wings", headers=h)).json()[0]["id"]
    await client.post(
        f"/api/v1/societies/{s.a_id}/flats", json={"wing_id": wing, "flat_number": "2"}, headers=h
    )
    done = (
        await client.post(f"/api/v1/societies/{s.a_id}/imports/{job['id']}/confirm", headers=h)
    ).json()
    assert done["ok_count"] == 1
    assert done["error_count"] == 1
    assert [r["status"] for r in done["rows"]] == ["ok", "error"]


async def test_not_an_excel_file(client, two_societies) -> None:  # noqa: ANN001
    s = two_societies
    r = await _upload(client, s.a_id, s.a_admin.headers, b"name,flat\nA,101\n")
    assert r.status_code == 422
    assert r.json()["error"]["code"] == "NOT_XLSX"


async def test_sheet_without_flat_number_column(client, two_societies) -> None:  # noqa: ANN001
    s = two_societies
    wb = Workbook()
    wb.active.append(["Name", "Phone"])
    buf = io.BytesIO()
    wb.save(buf)
    r = await _upload(client, s.a_id, s.a_admin.headers, buf.getvalue())
    assert r.json()["error"]["code"] == "TEMPLATE_COLUMNS"
