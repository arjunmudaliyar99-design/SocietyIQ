"""M6 acceptance (docs/07): the docs/06 §4 worked example end to end; no receipt for an
unconfirmed payment; reversal restores the balance exactly; every money action has an audit
row."""

import io

import asyncpg
from openpyxl import load_workbook

from app.auth import dev_user_id
from app.seed.demo import DEMO_ADMIN_EMAIL, run
from tests.conftest import APP_URL, TestUser, needs_db, new_user

pytestmark = needs_db

DEMO = TestUser(id=dev_user_id(DEMO_ADMIN_EMAIL), email=DEMO_ADMIN_EMAIL)
MONTHS = [f"2026-{m:02d}-01" for m in range(4, 11)]  # Apr..Oct: 7 bills


async def _sql(society_id: str, query: str, *args: object) -> list[asyncpg.Record]:
    conn = await asyncpg.connect(APP_URL)
    try:
        async with conn.transaction():
            await conn.execute("SELECT set_config('app.current_society_id', $1, true)", society_id)
            return await conn.fetch(query, *args)
    finally:
        await conn.close()


async def _seven_bills(client, admin) -> tuple[str, str, str]:  # noqa: ANN001
    """A society with one flat and ₹4,200 a month, billed Apr-Oct 2026 (₹29,400 owed)."""
    r = await client.post(
        "/api/v1/societies",
        json={"details": {"name": "Shanti Sadan CHS"}, "billing": {"bill_day": 1, "due_days": 15}},
        headers=admin.headers,
    )
    sid = r.json()["society"]["id"]
    base = f"/api/v1/societies/{sid}"
    h = admin.headers
    wing = (await client.get(f"{base}/wings", headers=h)).json()[0]["id"]
    flat = (
        await client.post(
            f"{base}/flats", json={"wing_id": wing, "flat_number": "B-207"}, headers=h
        )
    ).json()
    await client.post(
        f"{base}/charges/heads",
        json={
            "name": "Maintenance",
            "standard_code": "service",
            "calc_type": "equal_per_flat",
            "period_total_paise": 420000,
        },
        headers=h,
    )
    for month in MONTHS:
        draft = (
            await client.post(f"{base}/billing-runs", json={"period_start": month}, headers=h)
        ).json()
        assert (
            await client.post(f"{base}/billing-runs/{draft['id']}/approve", headers=h)
        ).status_code == 200
    return sid, base, flat["id"]


async def _balance(client, base: str, flat_id: str, h: dict) -> int:  # noqa: ANN001
    return (await client.get(f"{base}/flats/{flat_id}/ledger", headers=h)).json()["balance_paise"]


async def _bill_statuses(client, base: str, flat_id: str, h: dict) -> list[str]:  # noqa: ANN001
    bills = (await client.get(f"{base}/bills", params={"flat_id": flat_id}, headers=h)).json()
    return [b["status"] for b in sorted(bills, key=lambda b: b["bill_number"])]


async def test_worked_example_then_reversal_restores_exactly(client, two_societies) -> None:  # noqa: ANN001
    s = two_societies
    h = s.a_admin.headers
    sid, base, flat_id = await _seven_bills(client, s.a_admin)
    assert await _balance(client, base, flat_id, h) == 2940000

    r = await client.post(
        f"{base}/payments",
        json={
            "flat_id": flat_id,
            "amount_paise": 1000000,
            "method": "cash",
            "received_on": "2026-11-05",
            "confirm_now": True,
        },
        headers=h,
    )
    assert r.status_code == 201, r.text
    payment = r.json()
    assert payment["status"] == "confirmed"
    assert payment["receipt_number"] == "R/2026-27/000001"
    # Oldest two bills paid (₹8,400), the third partly (₹1,600); ₹19,400 left.
    assert [a["amount_paise"] for a in payment["allocations"]] == [420000, 420000, 160000]
    assert [a["settles"] for a in payment["allocations"]] == [
        "B/2026-27/000001",
        "B/2026-27/000002",
        "B/2026-27/000003",
    ]
    assert (
        await _bill_statuses(client, base, flat_id, h)
        == ["paid", "paid", "partially_paid"] + ["issued"] * 4
    )
    assert await _balance(client, base, flat_id, h) == 1940000
    dues = (await client.get(f"{base}/dues", headers=h)).json()
    assert dues["rows"][0]["months_pending"] == 5

    rev = await client.post(
        f"{base}/payments/{payment['id']}/reverse", json={"reason": "Counted twice"}, headers=h
    )
    assert rev.status_code == 200
    assert rev.json()["status"] == "reversed"
    assert rev.json()["receipt_cancelled"] is True
    assert rev.json()["allocations"] == []
    assert await _balance(client, base, flat_id, h) == 2940000
    assert await _bill_statuses(client, base, flat_id, h) == ["issued"] * 7
    ledger = (await client.get(f"{base}/flats/{flat_id}/ledger", headers=h)).json()
    assert [line["entry_type"] for line in ledger["lines"][-2:]] == ["payment", "reversal"]
    assert ledger["lines"][-1]["reason"] == "Counted twice"

    again = await client.post(
        f"{base}/payments/{payment['id']}/reverse", json={"reason": "again"}, headers=h
    )
    assert again.json()["error"]["code"] == "PAYMENT_NOT_REVERSIBLE"

    actions = [
        r["action"]
        for r in await _sql(
            sid,
            "SELECT action FROM audit_log WHERE entity_id = $1::uuid ORDER BY id",
            payment["id"],
        )
    ]
    assert actions == ["payment.record", "payment.confirm", "payment.reverse"]


async def test_no_receipt_until_confirmed(client, two_societies) -> None:  # noqa: ANN001
    s = two_societies
    h = s.a_admin.headers
    sid, base, flat_id = await _seven_bills(client, s.a_admin)
    upi = (
        await client.post(
            f"{base}/payments",
            json={
                "flat_id": flat_id,
                "amount_paise": 420000,
                "method": "upi",
                "received_on": "2026-11-02",
                "reference": "UTR123",
            },
            headers=h,
        )
    ).json()
    assert upi["status"] == "pending" and upi["receipt_number"] is None
    assert (await client.get(f"{base}/payments/{upi['id']}/receipt", headers=h)).json()["error"][
        "code"
    ] == "NO_RECEIPT"
    assert await _balance(client, base, flat_id, h) == 2940000  # no ledger effect yet

    confirmed = (await client.post(f"{base}/payments/{upi['id']}/confirm", headers=h)).json()
    assert confirmed["receipt_number"] == "R/2026-27/000001"
    assert await _balance(client, base, flat_id, h) == 2520000
    # Database-wide: every receipt belongs to a confirmed or later-reversed payment.
    rows = await _sql(sid, "SELECT p.status FROM receipts r JOIN payments p ON p.id = r.payment_id")
    assert {r["status"] for r in rows} <= {"confirmed", "reversed"}


async def test_cheque_rules_and_bounce(client, two_societies) -> None:  # noqa: ANN001
    s = two_societies
    h = s.a_admin.headers
    sid, base, flat_id = await _seven_bills(client, s.a_admin)
    body = {
        "flat_id": flat_id,
        "amount_paise": 420000,
        "method": "cheque",
        "received_on": "2026-11-03",
    }
    assert (await client.post(f"{base}/payments", json=body, headers=h)).json()["error"][
        "code"
    ] == "CHEQUE_NUMBER_REQUIRED"
    cheque = {**body, "reference": "000123", "cheque_bank": "Example Bank"}
    assert (
        await client.post(f"{base}/payments", json={**cheque, "confirm_now": True}, headers=h)
    ).json()["error"]["code"] == "CHEQUE_NOT_CLEARED"
    recorded = (await client.post(f"{base}/payments", json=cheque, headers=h)).json()
    assert recorded["status"] == "recorded"
    bounced = (
        await client.post(
            f"{base}/payments/{recorded['id']}/bounce",
            json={"reason": "Insufficient funds"},
            headers=h,
        )
    ).json()
    assert bounced["status"] == "bounced"
    assert await _balance(client, base, flat_id, h) == 2940000
    assert (await client.post(f"{base}/payments/{recorded['id']}/confirm", headers=h)).json()[
        "error"
    ]["code"] == "PAYMENT_NOT_OPEN"
    actions = [
        r["action"]
        for r in await _sql(
            sid,
            "SELECT action FROM audit_log WHERE entity_id = $1::uuid ORDER BY id",
            recorded["id"],
        )
    ]
    assert actions == ["payment.record", "payment.bounce"]


async def test_advance_settles_the_next_bill(client, two_societies) -> None:  # noqa: ANN001
    s = two_societies
    h = s.a_admin.headers
    sid, base, flat_id = await _seven_bills(client, s.a_admin)
    big = (
        await client.post(
            f"{base}/payments",
            json={
                "flat_id": flat_id,
                "amount_paise": 3400000,
                "method": "neft",
                "received_on": "2026-10-20",
                "reference": "N1",
                "confirm_now": True,
            },
            headers=h,
        )
    ).json()
    assert sum(a["amount_paise"] for a in big["allocations"]) == 2940000  # ₹4,600 advance
    assert await _balance(client, base, flat_id, h) == -460000
    nov = (
        await client.post(f"{base}/billing-runs", json={"period_start": "2026-11-01"}, headers=h)
    ).json()
    await client.post(f"{base}/billing-runs/{nov['id']}/approve", headers=h)
    statuses = await _bill_statuses(client, base, flat_id, h)
    assert statuses == ["paid"] * 8
    assert await _balance(client, base, flat_id, h) == -40000
    bill = next(
        b
        for b in (await client.get(f"{base}/bills", params={"flat_id": flat_id}, headers=h)).json()
        if b["bill_number"].endswith("08")
    )
    assert (bill["advance_paise"], bill["total_due_paise"]) == (460000, 0)


async def test_seed_dues_buckets_and_ledger_export(client) -> None:  # noqa: ANN001
    await run()
    mine = (await client.get("/api/v1/me/societies", headers=DEMO.headers)).json()
    sid = next(s["id"] for s in mine if s["name"] == "Gulmohar Heights CHS")
    base = f"/api/v1/societies/{sid}"
    dues = (await client.get(f"{base}/dues", headers=DEMO.headers)).json()
    by_flat = {
        r["flat_label"]: (r["balance_paise"], r["months_pending"], r["bucket"])
        for r in dues["rows"]
    }
    assert by_flat["3-04"] == (2520000, 7, "long")
    assert by_flat["2-02"] == (1080000, 3, "formal")
    assert by_flat["1-04"] == (360000, 1, "friendly")
    assert dues["counts"]["long"] >= 1
    only_long = (
        await client.get(f"{base}/dues", params={"bucket": "long"}, headers=DEMO.headers)
    ).json()
    assert {r["bucket"] for r in only_long["rows"]} == {"long"}

    flat_id = next(r["flat_id"] for r in dues["rows"] if r["flat_label"] == "3-04")
    xlsx = await client.get(
        f"{base}/flats/{flat_id}/ledger/export", params={"format": "xlsx"}, headers=DEMO.headers
    )
    assert xlsx.status_code == 200
    sheet = load_workbook(io.BytesIO(xlsx.content))["Ledger"]
    values = [c.value for c in sheet[5]]
    assert values[1] == "opening_balance" and values[6] == 25200
    resolutions = [row[0].value for row in sheet.iter_rows() if row[0].value == "Resolutions"]
    assert resolutions == ["Resolutions"]


async def test_resident_reports_payment_for_own_flat_only(client, two_societies) -> None:  # noqa: ANN001
    s = two_societies
    h = s.a_admin.headers
    sid, base, flat_id = await _seven_bills(client, s.a_admin)
    owner, tenant = new_user("owner"), new_user("tenant")
    for user, role in ((owner, "owner"), (tenant, "tenant")):
        inv = (
            await client.post(
                f"{base}/invites",
                json={"role": role, "email": user.email, "flat_id": flat_id},
                headers=h,
            )
        ).json()
        token = inv["invite_link"].split("token=", 1)[1]
        assert (
            await client.post(f"{base}/invites/accept", json={"token": token}, headers=user.headers)
        ).status_code == 200

    report = {
        "amount_paise": 420000,
        "method": "upi",
        "received_on": "2026-11-04",
        "reference": "UPI-REF-1",
    }
    r = await client.post(
        f"{base}/flats/{flat_id}/payment-reports", json=report, headers=owner.headers
    )
    assert r.status_code == 201, r.text
    assert (r.json()["status"], r.json()["reported_by_resident"]) == ("pending", True)
    assert await _balance(client, base, flat_id, h) == 2940000

    tenant_try = await client.post(
        f"{base}/flats/{flat_id}/payment-reports", json=report, headers=tenant.headers
    )
    assert tenant_try.json()["error"]["code"] == "TENANT_CANNOT_PAY"
    await client.post(
        f"{base}/flats/{flat_id}/members",
        json={
            "full_name": "Tenant B-207",
            "member_type": "tenant",
            "email": tenant.email,
            "tenant_can_pay": True,
        },
        headers=h,
    )
    assert (
        await client.post(
            f"{base}/flats/{flat_id}/payment-reports", json=report, headers=tenant.headers
        )
    ).status_code == 201

    wing = (await client.get(f"{base}/wings", headers=h)).json()[0]["id"]
    other = (
        await client.post(f"{base}/flats", json={"wing_id": wing, "flat_number": "X-1"}, headers=h)
    ).json()["id"]
    assert (
        await client.post(
            f"{base}/flats/{other}/payment-reports", json=report, headers=owner.headers
        )
    ).status_code == 404
    # Residents cannot see the committee's payment list or dues.
    assert (await client.get(f"{base}/payments", headers=owner.headers)).status_code == 403
    assert (await client.get(f"{base}/dues", headers=owner.headers)).status_code == 403


async def test_roles(client, two_societies) -> None:  # noqa: ANN001
    from tests.test_flats import _join

    s = two_societies
    sid, base, flat_id = await _seven_bills(client, s.a_admin)
    committee = await _join(client, sid, s.a_admin, "committee")
    auditor = await _join(client, sid, s.a_admin, "auditor")
    assert (await client.get(f"{base}/dues", headers=committee.headers)).status_code == 200
    assert (await client.get(f"{base}/payments", headers=committee.headers)).status_code == 403
    assert (await client.get(f"{base}/payments", headers=auditor.headers)).status_code == 200
    body = {"flat_id": flat_id, "amount_paise": 100, "method": "cash", "received_on": "2026-11-01"}
    assert (
        await client.post(f"{base}/payments", json=body, headers=auditor.headers)
    ).status_code == 403
