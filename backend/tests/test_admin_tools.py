"""M11: audit log viewer and data export (docs/07 M11, docs/02 §2)."""

import io
from decimal import Decimal

import asyncpg
from openpyxl import load_workbook

from app.auth import dev_user_id
from app.seed.demo import DEMO_ADMIN_EMAIL, run
from app.services.exports import SECRET_COLUMNS, society_tables
from tests.conftest import OWNER_URL, TestUser, needs_db

pytestmark = needs_db

DEMO = TestUser(id=dev_user_id(DEMO_ADMIN_EMAIL), email=DEMO_ADMIN_EMAIL)


async def test_every_table_with_society_data_is_exported() -> None:
    """A table added by a migration but not mapped as a model would be missing from exports."""
    conn = await asyncpg.connect(OWNER_URL)
    try:
        rows = await conn.fetch(
            "SELECT table_name FROM information_schema.columns"
            " WHERE table_schema = 'public' AND column_name = 'society_id'"
        )
    finally:
        await conn.close()
    assert {r["table_name"] for r in rows} == {t.name for t in society_tables()}


async def _shanti(client) -> str:  # noqa: ANN001
    await run()
    mine = (await client.get("/api/v1/me/societies", headers=DEMO.headers)).json()
    return next(s["id"] for s in mine if s["name"] == "Shanti Sadan CHS")


async def test_excel_export_has_all_tables_exact_rupees_and_no_secrets(client) -> None:  # noqa: ANN001
    sid = await _shanti(client)
    r = await client.get(f"/api/v1/societies/{sid}/export?format=xlsx", headers=DEMO.headers)
    assert r.status_code == 200, r.text
    assert "attachment" in r.headers["content-disposition"]
    book = load_workbook(io.BytesIO(r.content))
    assert set(book.sheetnames) == {t.name[:31] for t in society_tables()}

    flats = list(book["flats"].iter_rows(values_only=True))
    assert len(flats) - 1 == 120

    ledger = list(book["ledger_entries"].iter_rows(values_only=True))
    head = list(ledger[0])
    kind, amount = head.index("entry_type"), head.index("amount (₹)")
    opening = [row[amount] for row in ledger[1:] if row[kind] == "opening_balance"]
    # docs/02 §12 listed dues for Shanti Sadan: ₹75,600, written as exact rupees.
    assert sum(Decimal(str(v)) for v in opening) == Decimal("75600")

    for sheet in book.worksheets:
        headers = next(sheet.iter_rows(max_row=1, values_only=True), ()) or ()
        assert not SECRET_COLUMNS.intersection(headers), sheet.title

    society_col = list(book["flats"].iter_rows(max_row=1, values_only=True))[0].index("society_id")
    assert {row[society_col] for row in flats[1:]} == {sid}

    page = (
        await client.get(
            f"/api/v1/societies/{sid}/audit-log?action=society.export", headers=DEMO.headers
        )
    ).json()
    assert page["entries"][0]["after"]["format"] == "xlsx"


async def test_pdf_export_or_a_clear_503(client, two_societies) -> None:  # noqa: ANN001
    s = two_societies
    r = await client.get(f"/api/v1/societies/{s.a_id}/export?format=pdf", headers=s.a_admin.headers)
    # WeasyPrint needs Pango, which Windows machines usually lack (docs/09).
    assert r.status_code in (200, 503), r.text
    if r.status_code == 200:
        assert r.content.startswith(b"%PDF")
    else:
        assert r.json()["error"]["code"] == "PDF_UNAVAILABLE"


async def test_export_never_contains_another_society(client, two_societies) -> None:  # noqa: ANN001
    s = two_societies
    await client.post(
        f"/api/v1/societies/{s.b_id}/vendors", json={"name": "Only In B"}, headers=s.b_admin.headers
    )
    r = await client.get(f"/api/v1/societies/{s.a_id}/export", headers=s.a_admin.headers)
    book = load_workbook(io.BytesIO(r.content))
    for sheet in book.worksheets:
        for row in sheet.iter_rows(min_row=2, values_only=True):
            assert str(s.b_id) not in {str(v) for v in row}, sheet.title
            assert "Only In B" not in row


async def test_roles(client, two_societies) -> None:  # noqa: ANN001
    from tests.test_flats import _join

    s = two_societies
    base = f"/api/v1/societies/{s.a_id}"
    committee = await _join(client, s.a_id, s.a_admin, "committee")
    auditor = await _join(client, s.a_id, s.a_admin, "auditor")
    treasurer = await _join(client, s.a_id, s.a_admin, "treasurer")
    for user in (auditor, treasurer):
        assert (await client.get(f"{base}/audit-log", headers=user.headers)).status_code == 200
        assert (await client.get(f"{base}/export", headers=user.headers)).status_code == 200
    assert (await client.get(f"{base}/audit-log", headers=committee.headers)).status_code == 403
    assert (await client.get(f"{base}/export", headers=committee.headers)).status_code == 403


async def test_audit_log_newest_first_paged_filtered_and_named(client, two_societies) -> None:  # noqa: ANN001
    s = two_societies
    base = f"/api/v1/societies/{s.a_id}"
    h = s.a_admin.headers
    for n in range(55):
        await client.post(f"{base}/vendors", json={"name": f"Vendor {n:02d}"}, headers=h)

    first = (await client.get(f"{base}/audit-log?action=vendor.create", headers=h)).json()
    assert len(first["entries"]) == 50
    assert first["entries"][0]["after"]["name"] == "Vendor 54"
    ids = [e["id"] for e in first["entries"]]
    assert ids == sorted(ids, reverse=True)
    second = (
        await client.get(
            f"{base}/audit-log?action=vendor.create&before_id={first['next_before_id']}", headers=h
        )
    ).json()
    assert [e["after"]["name"] for e in second["entries"]] == [
        f"Vendor {n:02d}" for n in range(4, -1, -1)
    ]
    assert second["next_before_id"] is None

    everything = (await client.get(f"{base}/audit-log?entity_type=society", headers=h)).json()
    assert {e["entity_type"] for e in everything["entries"]} == {"society"}

    # Society B's admin sees none of A's rows.
    b_rows = (
        await client.get(
            f"/api/v1/societies/{s.b_id}/audit-log?action=vendor.create", headers=s.b_admin.headers
        )
    ).json()
    assert b_rows["entries"] == []
